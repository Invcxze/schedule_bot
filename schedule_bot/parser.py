"""Read-only parser for the supplied three-row timetable grid.

Merge resolution is coordinate-aware: an empty unmerged cell is NEVER filled
from its neighbour. Subject/teacher/room merges are resolved independently.
"""

import io
import re
from dataclasses import replace
from datetime import date, datetime
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .config import ParserConfig
from .models import GroupSchedule, Lesson, clean, normalize_group

MAX_XLSX_BYTES = 20 * 1024 * 1024
MAX_UNPACKED_BYTES = 100 * 1024 * 1024
TIME = re.compile(r"^(\d{1,2})[:.](\d{2})\s*[-–—−]\s*(\d{1,2})[:.](\d{2})$")
DAY_NAMES = {
    name: day
    for day, names in enumerate(
        (
            ("monday", "понедельник", "пн"),
            ("tuesday", "вторник", "вт"),
            ("wednesday", "среда", "ср"),
            ("thursday", "четверг", "чт"),
            ("friday", "пятница", "пт"),
            ("saturday", "суббота", "сб"),
            ("sunday", "воскресенье", "вс"),
        )
    )
    for name in names
}


class ParseError(ValueError):
    pass


def parse_day(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value).lower().rstrip(".")
    if text in DAY_NAMES:
        return f"w:{DAY_NAMES[text]}"
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def parse_time(value: object) -> tuple[str, str] | None:
    match = TIME.fullmatch(clean(value))
    if not match:
        return None
    h1, m1, h2, m2 = map(int, match.groups())
    if h1 > 23 or h2 > 23 or m1 > 59 or m2 > 59 or (h1, m1) >= (h2, m2):
        raise ParseError(f"Некорректный интервал времени: {value}")
    return f"{h1:02}:{m1:02}", f"{h2:02}:{m2:02}"


class Grid:
    def __init__(self, sheet):
        self.sheet = sheet
        self.anchors: dict[tuple[int, int], tuple[int, int]] = {}
        for merged in sheet.merged_cells.ranges:
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    self.anchors[row, col] = (merged.min_row, merged.min_col)

    def anchor(self, row: int, col: int) -> tuple[int, int]:
        return self.anchors.get((row, col), (row, col))

    def value(self, row: int, col: int) -> object:
        return self.sheet.cell(*self.anchor(row, col)).value


def _parse_sheet(sheet, name: str, config: ParserConfig) -> dict[str, GroupSchedule]:
    if sheet.max_row > 10000 or sheet.max_column > 300:
        raise ParseError(f"Слишком большая сетка на листе {name!r}.")
    grid = Grid(sheet)
    groups = []
    for cell in sheet[config.header_row]:
        group = normalize_group(clean(cell.value))
        if re.fullmatch(config.group_pattern, group):
            groups.append((cell.column, group))
    if not groups:
        raise ParseError(f"На листе {name!r} не найдены группы в строке {config.header_row}.")
    # A year block has its own time column, including different start times.
    time_columns = {
        col
        for col in range(1, sheet.max_column + 1)
        if any(
            parse_time(sheet.cell(row, col).value)
            for row in range(config.header_row + 1, sheet.max_row + 1)
        )
    }
    schedules: dict[str, GroupSchedule] = {}
    for group_col, group in groups:
        if group in schedules:
            raise ParseError(
                f"Группа {group} встречается дважды на листе {name!r}. "
                "Выбери одну актуальную версию."
            )
        candidates = [col for col in time_columns if col < group_col]
        if not candidates:
            raise ParseError(f"Не найдена колонка времени для {group} на {name!r}.")
        time_col = max(candidates)
        days: dict[str, list[Lesson]] = {}
        current_day = None
        slots = 0
        for row in range(config.header_row + 1, sheet.max_row + 1):
            # Only raw time anchors start a slot; vertical merges are not extra lessons.
            raw_time = sheet.cell(row, time_col).value
            day = parse_day(raw_time)
            if day:
                if day in days:
                    raise ParseError(f"Повторяется день {day} для {group} на {name!r}.")
                current_day = day
                days[day] = []
                continue
            interval = parse_time(raw_time)
            if interval is None:
                if clean(raw_time):
                    raise ParseError(f"Неизвестная метка дня/времени {name}!{row}: {raw_time}")
                continue
            if current_day is None:
                raise ParseError(f"Время без дня на листе {name}, строка {row}.")
            slots += 1
            if any(
                clean(sheet.cell(r, time_col).value)
                for r in range(row + 1, row + config.lesson_rows)
            ):
                raise ParseError(f"Нарушена трёхстрочная структура: {name}, строка {row}.")
            subject = clean(grid.value(row, group_col))
            if not subject or subject in {"-", "—", "–"}:
                continue
            # A subject may cover all 3 rows (e.g. Physical Education).
            seen = {grid.anchor(row, group_col)}
            details = []
            for offset in (1, 2):
                anchor = grid.anchor(row + offset, group_col)
                details.append(
                    "" if anchor in seen else clean(grid.value(row + offset, group_col))
                )
                seen.add(anchor)
            days[current_day].append(Lesson(*interval, subject, *details))
        if not days or not slots:
            raise ParseError(f"Нет распознаваемой сетки расписания для {group}.")
        # Sunday is implicit in the supplied six-day weekly grid.
        if any(day.startswith("w:") for day in days):
            days.setdefault("w:6", [])
        schedules[group] = GroupSchedule(
            group,
            (name,),
            {d: tuple(sorted(set(items))) for d, items in days.items()},
            config.valid_from,
            config.valid_until,
        )
    return schedules


def _merge(base: GroupSchedule, extra: GroupSchedule) -> GroupSchedule:
    """Union two partial schedules for the same group from different sheets.

    Some groups are split across sheets (e.g. common courses plus a
    program-specific sheet with extra electives); identical lessons listed on
    both sheets collapse via the same Lesson-set dedup used within one sheet.
    """
    days = {
        key: tuple(sorted(set(base.days.get(key, ())) | set(extra.days.get(key, ()))))
        for key in set(base.days) | set(extra.days)
    }
    sheets = base.sheets + tuple(s for s in extra.sheets if s not in base.sheets)
    return replace(base, sheets=sheets, days=days)


def parse_workbook(data: bytes, config: ParserConfig) -> dict[str, GroupSchedule]:
    if len(data) > MAX_XLSX_BYTES:
        raise ParseError("XLSX больше 20 МБ.")
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise ParseError("Распакованный XLSX больше 100 МБ.")
        workbook = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ParseError(
            "Источник не является корректным XLSX (возможно, страница входа Google)."
        ) from exc
    schedules: dict[str, GroupSchedule] = {}
    try:
        for name in config.sheet_names:
            if name not in workbook.sheetnames:
                raise ParseError(
                    f"Не найден лист {name!r}. Доступны: {', '.join(workbook.sheetnames)}"
                )
            for group, schedule in _parse_sheet(workbook[name], name, config).items():
                # A group split across configured sheets is merged, not rejected: the
                # same physical group commonly has core courses on one sheet and
                # program-specific electives on another (see _merge).
                schedules[group] = (
                    _merge(schedules[group], schedule) if group in schedules else schedule
                )
    finally:
        workbook.close()
    return schedules
